"""Compressed workbook size cannot bypass the upload budget."""

import re
from io import BytesIO
from unittest.mock import Mock
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from starlette.datastructures import UploadFile

from app.services import excel_import_export as excel


def _workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Personnel"
    ws.append(["name", "role", "status"])
    ws.append(["Test Person", "", "available"])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def test_expanded_workbook_limit_is_checked_before_openpyxl(monkeypatch):
    # A small configured budget exercises the real compressed ZIP expansion without
    # asking the regression suite to allocate a production-sized decompression bomb.
    monkeypatch.setattr(excel, "MAX_WORKBOOK_UNCOMPRESSED_BYTES", 1024, raising=False)
    parser = Mock(wraps=excel.openpyxl.load_workbook)
    monkeypatch.setattr(excel.openpyxl, "load_workbook", parser)

    with pytest.raises(excel.ExcelImportError, match="entpackt"):
        excel.validate_and_parse_excel(_workbook_bytes())
    parser.assert_not_called()


def test_declared_sheet_dimensions_cannot_create_unbounded_iteration():
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:C2"', b'ref="A1:C1048576"')
            target.writestr(entry.filename, content)

    with pytest.raises(excel.ExcelImportError, match="Zeilen"):
        excel.validate_and_parse_excel(output.getvalue())


@pytest.mark.parametrize("dimension", [b"A1:C1", b"A1:A2", b"A1:A1"])
def test_underdeclared_dimensions_do_not_discard_roster_data(dimension):
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:C2"', b'ref="' + dimension + b'"')
            target.writestr(entry.filename, content)

    parsed = excel.validate_and_parse_excel(output.getvalue())
    assert parsed.personnel.rows == [{"name": "Test Person", "role": None, "status": "available"}]


@pytest.mark.parametrize("column", [b"BM", b"BN", b"ZZZZZZZ"])
def test_actual_cell_coordinates_are_bounded_without_dimensions(column):
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb"<dimension[^>]+/>", b"", content)
                cell = b'<c r="' + column + b'2" t="inlineStr"><is><t>Extra</t></is></c>'
                content = content.replace(b"</row></sheetData>", cell + b"</row></sheetData>")
            target.writestr(entry.filename, content)

    with pytest.raises(excel.ExcelImportError, match="Spalten"):
        excel.validate_and_parse_excel(output.getvalue())


@pytest.mark.parametrize("part", ["xl/custom-sheet.xml", "xl/custom-sheet.part"])
def test_coordinate_scan_follows_nonstandard_sheet_locations(monkeypatch, part):
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            name = entry.filename
            if name == "xl/worksheets/sheet1.xml":
                name = part
                cell = b'<c r="BM2" t="inlineStr"><is><t>Extra</t></is></c>'
                content = content.replace(b"</row></sheetData>", cell + b"</row></sheetData>")
            elif name in ("[Content_Types].xml", "xl/_rels/workbook.xml.rels"):
                content = content.replace(b"/xl/worksheets/sheet1.xml", b"/" + part.encode())
            target.writestr(name, content)
    parser = Mock(wraps=excel.openpyxl.load_workbook)
    monkeypatch.setattr(excel.openpyxl, "load_workbook", parser)

    with pytest.raises(excel.ExcelImportError, match="Spalten"):
        excel.validate_and_parse_excel(output.getvalue())
    parser.assert_not_called()


@pytest.mark.parametrize("root_tag,cell_tag", [(b"customroot", b"c"), (b"worksheet", b"unexpected")])
def test_nonstandard_xml_cannot_hide_unchecked_cell_coordinates(monkeypatch, root_tag, cell_tag):
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"worksheet", root_tag)
                cell = b"<" + cell_tag + b' r="BM2" t="inlineStr"><is><t>Extra</t></is></' + cell_tag + b">"
                content = content.replace(b"</row></sheetData>", cell + b"</row></sheetData>")
            target.writestr(entry.filename, content)
    parser = Mock(wraps=excel.openpyxl.load_workbook)
    monkeypatch.setattr(excel.openpyxl, "load_workbook", parser)

    with pytest.raises(excel.ExcelImportError):
        excel.validate_and_parse_excel(output.getvalue())
    parser.assert_not_called()


def test_valid_row_extensions_and_unrelated_parts_remain_importable():
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"</row></sheetData>", b"<extLst/></row></sheetData>")
            target.writestr(entry.filename, content)
        target.writestr("xl/media/example.bin", b"\x89PNG\x00not sheet data")
        target.writestr(
            "xl/calcChain.xml",
            b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            + b'<c r="A1"/>' * 100
            + b"</calcChain>",
        )

    parsed = excel.validate_and_parse_excel(output.getvalue())
    assert parsed.personnel.rows == [{"name": "Test Person", "role": None, "status": "available"}]


def test_xml_entities_are_rejected():
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = b'<!DOCTYPE worksheet [<!ENTITY injected "Entity Person">]>' + content
                content = content.replace(b"Test Person", b"&injected;")
            target.writestr(entry.filename, content)

    with pytest.raises(excel.ExcelImportError):
        excel.validate_and_parse_excel(output.getvalue())


def test_row_budget_also_applies_without_declared_dimensions(monkeypatch):
    monkeypatch.setattr(excel, "MAX_IMPORT_ROWS", 2)
    output = BytesIO()
    with ZipFile(BytesIO(_workbook_bytes())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb"<dimension[^>]+/>", b"", content)
                row = b'<row r="3"><c r="A3" t="inlineStr"><is><t>Second</t></is></c></row>'
                row += b'<row r="4"><c r="A4" t="inlineStr"><is><t>Third</t></is></c></row>'
                content = content.replace(b"</sheetData>", row + b"</sheetData>")
            target.writestr(entry.filename, content)

    with pytest.raises(excel.ExcelImportError, match="Zeilen"):
        excel.validate_and_parse_excel(output.getvalue())


@pytest.mark.asyncio
@pytest.mark.parametrize("endpoint", ["preview", "execute"])
async def test_import_upload_is_read_with_a_byte_budget(admin_client, monkeypatch, endpoint):
    from app.api.admin import settings

    monkeypatch.setattr(settings, "max_excel_import_mb", 1)
    read_sizes = []
    original_read = UploadFile.read

    async def read(upload, size=-1):
        read_sizes.append(size)
        return await original_read(upload, size)

    monkeypatch.setattr(UploadFile, "read", read)
    response = await admin_client.post(
        f"/api/admin/import/{endpoint}",
        files={"file": ("roster.xlsx", _workbook_bytes())},
        data={"mode": "append"},
    )

    assert response.status_code == 200
    assert read_sizes == [1024 * 1024 + 1]
